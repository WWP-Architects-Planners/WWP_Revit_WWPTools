#!python3
# -*- coding: utf-8 -*-

import os
from pyrevit import DB
from System import Activator, Type
from System.Runtime.InteropServices import Marshal

import WWP_colorSchemeUtils as csu
try:
    from openpyxl import Workbook, load_workbook
except Exception:
    Workbook = None
    load_workbook = None

try:
    Reflection = __import__('System.Reflection', fromlist=['BindingFlags'])
    BindingFlags = Reflection.BindingFlags
except Exception:
    BindingFlags = None


EXCEL_SHEET_NAME = "ColorScheme"
FORMAT_VERSION = "1"
_BASE_FLAGS = (BindingFlags.Public | BindingFlags.Instance | BindingFlags.OptionalParamBinding) if BindingFlags is not None else None


def elem_id_int(elem_id):
    try:
        return int(elem_id.IntegerValue)
    except Exception:
        try:
            return int(elem_id)
        except Exception:
            return None


def scheme_area_scheme_id(scheme):
    try:
        area_scheme_id = getattr(scheme, "AreaSchemeId", None)
        if area_scheme_id and area_scheme_id != DB.ElementId.InvalidElementId:
            return area_scheme_id
    except Exception:
        pass
    try:
        getter = getattr(scheme, "GetAreaSchemeId", None)
        if callable(getter):
            area_scheme_id = getter()
            if area_scheme_id and area_scheme_id != DB.ElementId.InvalidElementId:
                return area_scheme_id
    except Exception:
        pass
    return None


def scheme_area_scheme_name(doc, scheme):
    try:
        area_scheme_id = scheme_area_scheme_id(scheme)
        if area_scheme_id:
            area_scheme = doc.GetElement(area_scheme_id)
            if area_scheme:
                return getattr(area_scheme, "Name", "") or ""
    except Exception:
        pass
    return ""


def category_name(doc, category_id):
    if category_id is None:
        return "Unknown Category"
    cat_int = elem_id_int(category_id)
    try:
        cat = doc.GetElement(category_id)
        cat_name = getattr(cat, "Name", "") if cat else ""
        if cat_name:
            return cat_name
    except Exception:
        pass
    try:
        categories = getattr(getattr(doc, "Settings", None), "Categories", None)
        if categories is not None:
            cat = categories.get_Item(category_id)
            cat_name = getattr(cat, "Name", "") if cat else ""
            if cat_name:
                return cat_name
    except Exception:
        pass
    try:
        import System
        bic = System.Enum.ToObject(DB.BuiltInCategory, cat_int)
        categories = getattr(getattr(doc, "Settings", None), "Categories", None)
        if categories is not None:
            cat = categories.get_Item(bic)
            cat_name = getattr(cat, "Name", "") if cat else ""
            if cat_name:
                return cat_name
    except Exception:
        pass
    try:
        import System
        bic = System.Enum.ToObject(DB.BuiltInCategory, cat_int)
        label = DB.LabelUtils.GetLabelFor(bic)
        if label:
            return label
    except Exception:
        pass
    try:
        label = DB.LabelUtils.GetLabelFor(category_id)
        if label:
            return label
    except Exception:
        pass
    if cat_int is not None:
        return "Category {}".format(cat_int)
    return "Unknown Category"


def scheme_scope_label(doc, scheme):
    area_name = scheme_area_scheme_name(doc, scheme)
    if area_name:
        return "Area({})".format(area_name)
    return category_name(doc, getattr(scheme, "CategoryId", None))


def scheme_display_name(doc, scheme):
    scheme_name = getattr(scheme, "Name", "") or "Color Scheme"
    return "{}: {}".format(scheme_scope_label(doc, scheme), scheme_name)


def describe_scheme(scheme):
    return "scheme='{}' id={} categoryId={} areaSchemeId={}".format(
        getattr(scheme, "Name", "<unnamed>"),
        elem_id_int(getattr(scheme, "Id", None)),
        elem_id_int(getattr(scheme, "CategoryId", None)),
        elem_id_int(scheme_area_scheme_id(scheme)),
    )


def storage_type_token(storage_type):
    if storage_type == DB.StorageType.String:
        return "String"
    if storage_type == DB.StorageType.Integer:
        return "Integer"
    if storage_type == DB.StorageType.Double:
        return "Double"
    if storage_type == DB.StorageType.ElementId:
        return "ElementId"
    return "Unknown"


def parse_storage_type(token):
    value = (token or "").strip().lower()
    if value == "string":
        return DB.StorageType.String
    if value == "integer":
        return DB.StorageType.Integer
    if value == "double":
        return DB.StorageType.Double
    if value == "elementid":
        return DB.StorageType.ElementId
    return None


def entry_value_to_text(entry, storage_type):
    value = csu._entry_get_value(entry, storage_type)
    if storage_type == DB.StorageType.Double:
        try:
            return "{:.12g}".format(float(value))
        except Exception:
            return ""
    if storage_type == DB.StorageType.ElementId:
        return "" if value is None else str(elem_id_int(value))
    return "" if value is None else str(value)


def _scheme_parameter_id(scheme):
    try:
        pid = getattr(scheme, "ParameterId", None)
        if pid is not None:
            return pid
    except Exception:
        pass
    return None


def _scheme_parameter_id_int(scheme):
    return elem_id_int(_scheme_parameter_id(scheme))


def _scheme_parameter_label(doc, scheme):
    pid = _scheme_parameter_id(scheme)
    if pid is None:
        return ""
    try:
        if pid == DB.ElementId.InvalidElementId:
            return ""
    except Exception:
        pass
    try:
        elem = doc.GetElement(pid)
        name = getattr(elem, "Name", "") if elem else ""
        if name:
            return name
    except Exception:
        pass
    try:
        binding = getattr(scheme, "ParameterDefinition", None)
        name = getattr(binding, "Name", "") if binding else ""
        if name:
            return name
    except Exception:
        pass
    pid_int = elem_id_int(pid)
    return "" if pid_int is None else "ParameterId {}".format(pid_int)


def build_payload_from_scheme(doc, scheme):
    payload = {
        "format_version": FORMAT_VERSION,
        "scheme_name": getattr(scheme, "Name", "") or "Color Scheme",
        "category_name": category_name(doc, getattr(scheme, "CategoryId", None)),
        "category_id": elem_id_int(getattr(scheme, "CategoryId", None)),
        "area_scheme_name": scheme_area_scheme_name(doc, scheme),
        "title": getattr(scheme, "Title", "") or "",
        "is_by_range": bool(getattr(scheme, "IsByRange", False)),
        "is_by_value": bool(getattr(scheme, "IsByValue", False)),
        "is_by_percentage": bool(getattr(scheme, "IsByPercentage", False)),
        "parameter_id": _scheme_parameter_id_int(scheme),
        "parameter_name": _scheme_parameter_label(doc, scheme),
        "entries": [],
    }
    try:
        entries = list(scheme.GetEntries())
    except Exception:
        entries = []
    for entry in entries:
        color = getattr(entry, "Color", None)
        payload["entries"].append({
            "caption": csu._entry_caption(entry),
            "storage_type": storage_type_token(getattr(entry, "StorageType", None)),
            "value": entry_value_to_text(entry, getattr(entry, "StorageType", None)),
            "color_r": "" if color is None else int(getattr(color, "Red", 0)),
            "color_g": "" if color is None else int(getattr(color, "Green", 0)),
            "color_b": "" if color is None else int(getattr(color, "Blue", 0)),
            "fill_pattern_id": elem_id_int(getattr(entry, "FillPatternId", None)),
            "is_visible": bool(getattr(entry, "IsVisible", getattr(entry, "Visible", True))),
        })
    return payload


def _com_get(obj, name):
    try:
        return obj.GetType().InvokeMember(name, _BASE_FLAGS | BindingFlags.GetProperty, None, obj, None)
    except Exception:
        return obj.GetType().InvokeMember(name, _BASE_FLAGS | BindingFlags.InvokeMethod, None, obj, [])


def _com_set(obj, name, value):
    try:
        obj.GetType().InvokeMember(name, _BASE_FLAGS | BindingFlags.SetProperty, None, obj, [value])
    except Exception:
        obj.GetType().InvokeMember(name, _BASE_FLAGS | BindingFlags.InvokeMethod, None, obj, [value])


def _com_call(obj, name, *args):
    try:
        return obj.GetType().InvokeMember(name, _BASE_FLAGS | BindingFlags.InvokeMethod, None, obj, list(args))
    except Exception:
        return obj.GetType().InvokeMember(name, _BASE_FLAGS | BindingFlags.GetProperty, None, obj, list(args))


def _com_call_strict(obj, name, *args):
    return obj.GetType().InvokeMember(name, _BASE_FLAGS | BindingFlags.InvokeMethod, None, obj, list(args))


def _try_delete_file(path, log=None):
    try:
        if os.path.isfile(path):
            os.remove(path)
            _log(log, "Deleted existing file before save: {}".format(path))
    except Exception as ex:
        _log(log, "Could not delete existing file '{}': {}".format(path, str(ex)))


def _save_workbook_with_fallbacks(workbook, path, log=None):
    # 51 = xlOpenXMLWorkbook (.xlsx)
    save_attempts = [
        ("SaveAs-xlsx", lambda: _com_call_strict(workbook, "SaveAs", path, 51)),
        ("SaveAs-basic", lambda: _com_call_strict(workbook, "SaveAs", path)),
        ("SaveCopyAs", lambda: _com_call_strict(workbook, "SaveCopyAs", path)),
    ]
    for label, action in save_attempts:
        _try_delete_file(path, log=log)
        _log(log, "Trying workbook save strategy: {}".format(label))
        try:
            action()
        except Exception as ex:
            _log(log, "Save strategy {} failed: {}".format(label, str(ex)))
            continue
        if os.path.isfile(path):
            _log(log, "Save strategy {} succeeded.".format(label))
            return True
        _log(log, "Save strategy {} returned without error, but file is still missing.".format(label))
    return False


def _log(log, message):
    if callable(log):
        try:
            log(message)
        except Exception:
            pass


def export_payload_to_excel(payload, path, log=None):
    if Workbook is not None:
        _log(log, "Using openpyxl export path.")
        wb = Workbook()
        ws = wb.active
        ws.title = EXCEL_SHEET_NAME

        rows = [
            ["WWP Color Scheme Export", ""],
            ["FormatVersion", payload.get("format_version", FORMAT_VERSION)],
            ["SchemeName", payload.get("scheme_name", "")],
            ["CategoryName", payload.get("category_name", "")],
            ["CategoryId", payload.get("category_id", "")],
            ["AreaSchemeName", payload.get("area_scheme_name", "")],
            ["Title", payload.get("title", "")],
            ["ParameterId", payload.get("parameter_id", "")],
            ["ParameterName", payload.get("parameter_name", "")],
            ["IsByRange", "TRUE" if payload.get("is_by_range") else "FALSE"],
            ["IsByValue", "TRUE" if payload.get("is_by_value") else "FALSE"],
            ["IsByPercentage", "TRUE" if payload.get("is_by_percentage") else "FALSE"],
            [],
            ["Caption", "StorageType", "Value", "ColorR", "ColorG", "ColorB", "FillPatternId", "IsVisible"],
        ]
        for entry in payload.get("entries", []):
            rows.append([
                entry.get("caption", ""),
                entry.get("storage_type", ""),
                entry.get("value", ""),
                entry.get("color_r", ""),
                entry.get("color_g", ""),
                entry.get("color_b", ""),
                entry.get("fill_pattern_id", ""),
                "TRUE" if entry.get("is_visible", True) else "FALSE",
            ])

        _log(log, "Prepared {} rows for openpyxl export.".format(len(rows)))
        for row in rows:
            ws.append(row)
        for column_cells in ws.columns:
            width = 0
            for cell in column_cells:
                try:
                    width = max(width, len("" if cell.value is None else str(cell.value)))
                except Exception:
                    pass
            col_letter = column_cells[0].column_letter
            ws.column_dimensions[col_letter].width = min(max(width + 2, 10), 60)

        directory = os.path.dirname(path)
        if directory and not os.path.isdir(directory):
            os.makedirs(directory)
            _log(log, "Created export directory: {}".format(directory))
        _try_delete_file(path, log=log)
        wb.save(path)
        if not os.path.isfile(path):
            raise Exception("openpyxl completed without creating '{}'.".format(path))
        _log(log, "Workbook saved successfully via openpyxl.")
        return

    _log(log, "openpyxl unavailable. Falling back to Excel COM export.")
    excel_app = None
    workbook = None
    worksheet = None
    used_range = None
    try:
        _log(log, "Starting Excel export.")
        excel_type = Type.GetTypeFromProgID("Excel.Application")
        if excel_type is None:
            raise Exception("Excel COM ProgID not found. Verify Microsoft Excel is installed.")
        _log(log, "Excel COM type acquired.")
        excel_app = Activator.CreateInstance(excel_type)
        _log(log, "Excel application created.")
        _com_set(excel_app, "Visible", False)
        _com_set(excel_app, "DisplayAlerts", False)
        workbooks = _com_get(excel_app, "Workbooks")
        _log(log, "Excel workbooks collection acquired.")
        workbook = _com_call(workbooks, "Add")
        _log(log, "Workbook created.")
        worksheets = _com_get(workbook, "Worksheets")
        worksheet = _com_call(worksheets, "Item", 1)
        _com_set(worksheet, "Name", EXCEL_SHEET_NAME)
        _log(log, "Worksheet prepared: {}".format(EXCEL_SHEET_NAME))

        rows = [
            ["WWP Color Scheme Export", ""],
            ["FormatVersion", payload.get("format_version", FORMAT_VERSION)],
            ["SchemeName", payload.get("scheme_name", "")],
            ["CategoryName", payload.get("category_name", "")],
            ["CategoryId", payload.get("category_id", "")],
            ["AreaSchemeName", payload.get("area_scheme_name", "")],
            ["Title", payload.get("title", "")],
            ["ParameterId", payload.get("parameter_id", "")],
            ["ParameterName", payload.get("parameter_name", "")],
            ["IsByRange", "TRUE" if payload.get("is_by_range") else "FALSE"],
            ["IsByValue", "TRUE" if payload.get("is_by_value") else "FALSE"],
            ["IsByPercentage", "TRUE" if payload.get("is_by_percentage") else "FALSE"],
            [],
            ["Caption", "StorageType", "Value", "ColorR", "ColorG", "ColorB", "FillPatternId", "IsVisible"],
        ]
        for entry in payload.get("entries", []):
            rows.append([
                entry.get("caption", ""),
                entry.get("storage_type", ""),
                entry.get("value", ""),
                entry.get("color_r", ""),
                entry.get("color_g", ""),
                entry.get("color_b", ""),
                entry.get("fill_pattern_id", ""),
                "TRUE" if entry.get("is_visible", True) else "FALSE",
            ])

        row_count = len(rows)
        col_count = max(len(r) for r in rows) if rows else 1
        _log(log, "Prepared {} rows x {} columns for export.".format(row_count, col_count))
        cells = _com_get(worksheet, "Cells")
        for r_idx, row in enumerate(rows, start=1):
            for c_idx in range(1, col_count + 1):
                value = row[c_idx - 1] if c_idx - 1 < len(row) else ""
                cell = _com_call(cells, "Item", r_idx, c_idx)
                _com_set(cell, "Value2", value)
        _log(log, "Cell values written to worksheet.")

        start_cell = _com_call(cells, "Item", 1, 1)
        end_cell = _com_call(cells, "Item", row_count, col_count)
        range_obj = _com_call(worksheet, "Range", start_cell, end_cell)
        used_range = _com_get(worksheet, "UsedRange")
        _com_call(_com_get(used_range, "Columns"), "AutoFit")
        _log(log, "Columns auto-fitted.")

        directory = os.path.dirname(path)
        if directory and not os.path.isdir(directory):
            os.makedirs(directory)
            _log(log, "Created export directory: {}".format(directory))
        _log(log, "Saving workbook to: {}".format(path))
        if not _save_workbook_with_fallbacks(workbook, path, log=log):
            raise Exception("Excel reported success but no file was created at '{}'.".format(path))
        _log(log, "Workbook saved successfully.")
    finally:
        try:
            if workbook is not None:
                _log(log, "Closing workbook.")
                _com_call_strict(workbook, "Close", False)
        except Exception:
            pass
        try:
            if excel_app is not None:
                _log(log, "Quitting Excel application.")
                _com_call_strict(excel_app, "Quit")
        except Exception:
            pass
        for obj in (used_range, worksheet, workbook, excel_app):
            try:
                if obj is not None:
                    Marshal.ReleaseComObject(obj)
            except Exception:
                pass


def _to_bool(text):
    return str(text or "").strip().lower() in ("true", "1", "yes", "y")


def _to_int(text, default=None):
    try:
        if text is None or str(text).strip() == "":
            return default
        return int(float(str(text).strip()))
    except Exception:
        return default


def _to_float(text, default=None):
    try:
        if text is None or str(text).strip() == "":
            return default
        return float(str(text).strip())
    except Exception:
        return default


def import_payload_from_excel(path):
    if load_workbook is not None:
        wb = load_workbook(path, data_only=True)
        if EXCEL_SHEET_NAME not in wb.sheetnames:
            raise Exception("Worksheet '{}' was not found in the workbook.".format(EXCEL_SHEET_NAME))
        ws = wb[EXCEL_SHEET_NAME]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if value is None else str(value) for value in row])

        if not rows or not rows[0] or rows[0][0].strip() != "WWP Color Scheme Export":
            raise Exception("The workbook does not match the WWP color scheme export format.")

        metadata = {}
        entry_start = None
        for idx, row in enumerate(rows[1:], start=1):
            first = row[0].strip() if row else ""
            if first == "Caption":
                entry_start = idx + 1
                break
            if first:
                metadata[first] = row[1].strip() if len(row) > 1 else ""
        if entry_start is None:
            raise Exception("Entry header row was not found in the workbook.")

        payload = {
            "format_version": metadata.get("FormatVersion", FORMAT_VERSION),
            "scheme_name": metadata.get("SchemeName", "Imported Color Scheme"),
            "category_name": metadata.get("CategoryName", ""),
            "category_id": _to_int(metadata.get("CategoryId"), None),
            "area_scheme_name": metadata.get("AreaSchemeName", ""),
            "title": metadata.get("Title", ""),
            "parameter_id": _to_int(metadata.get("ParameterId"), None),
            "parameter_name": metadata.get("ParameterName", ""),
            "is_by_range": _to_bool(metadata.get("IsByRange")),
            "is_by_value": _to_bool(metadata.get("IsByValue")),
            "is_by_percentage": _to_bool(metadata.get("IsByPercentage")),
            "entries": [],
        }
        for row in rows[entry_start:]:
            if not any((cell or "").strip() for cell in row):
                continue
            payload["entries"].append({
                "caption": row[0].strip() if len(row) > 0 else "",
                "storage_type": row[1].strip() if len(row) > 1 else "",
                "value": row[2].strip() if len(row) > 2 else "",
                "color_r": _to_int(row[3] if len(row) > 3 else "", None),
                "color_g": _to_int(row[4] if len(row) > 4 else "", None),
                "color_b": _to_int(row[5] if len(row) > 5 else "", None),
                "fill_pattern_id": _to_int(row[6] if len(row) > 6 else "", None),
                "is_visible": _to_bool(row[7] if len(row) > 7 else "TRUE"),
            })
        return payload

    excel_app = None
    workbook = None
    worksheet = None
    used_range = None
    try:
        excel_type = Type.GetTypeFromProgID("Excel.Application")
        if excel_type is None:
            raise Exception("Excel COM ProgID not found. Verify Microsoft Excel is installed.")
        excel_app = Activator.CreateInstance(excel_type)
        _com_set(excel_app, "Visible", False)
        _com_set(excel_app, "DisplayAlerts", False)
        workbooks = _com_get(excel_app, "Workbooks")
        workbook = _com_call(workbooks, "Open", path)
        worksheets = _com_get(workbook, "Worksheets")
        worksheet = _com_call(worksheets, "Item", EXCEL_SHEET_NAME)
        used_range = _com_get(worksheet, "UsedRange")
        values = _com_get(used_range, "Value2")
        rows_count = int(_com_get(_com_get(used_range, "Rows"), "Count"))
        cols_count = int(_com_get(_com_get(used_range, "Columns"), "Count"))
        lower_r = int(values.GetLowerBound(0))
        lower_c = int(values.GetLowerBound(1))

        rows = []
        for r in range(rows_count):
            row = []
            for c in range(cols_count):
                value = values.GetValue(lower_r + r, lower_c + c)
                row.append("" if value is None else str(value))
            rows.append(row)

        if not rows or not rows[0] or rows[0][0].strip() != "WWP Color Scheme Export":
            raise Exception("The workbook does not match the WWP color scheme export format.")

        metadata = {}
        entry_start = None
        for idx, row in enumerate(rows[1:], start=1):
            first = row[0].strip() if row else ""
            if first == "Caption":
                entry_start = idx + 1
                break
            if first:
                metadata[first] = row[1].strip() if len(row) > 1 else ""
        if entry_start is None:
            raise Exception("Entry header row was not found in the workbook.")

        payload = {
            "format_version": metadata.get("FormatVersion", FORMAT_VERSION),
            "scheme_name": metadata.get("SchemeName", "Imported Color Scheme"),
            "category_name": metadata.get("CategoryName", ""),
            "category_id": _to_int(metadata.get("CategoryId"), None),
            "area_scheme_name": metadata.get("AreaSchemeName", ""),
            "title": metadata.get("Title", ""),
            "parameter_id": _to_int(metadata.get("ParameterId"), None),
            "parameter_name": metadata.get("ParameterName", ""),
            "is_by_range": _to_bool(metadata.get("IsByRange")),
            "is_by_value": _to_bool(metadata.get("IsByValue")),
            "is_by_percentage": _to_bool(metadata.get("IsByPercentage")),
            "entries": [],
        }
        for row in rows[entry_start:]:
            if not any((cell or "").strip() for cell in row):
                continue
            payload["entries"].append({
                "caption": row[0].strip() if len(row) > 0 else "",
                "storage_type": row[1].strip() if len(row) > 1 else "",
                "value": row[2].strip() if len(row) > 2 else "",
                "color_r": _to_int(row[3] if len(row) > 3 else "", None),
                "color_g": _to_int(row[4] if len(row) > 4 else "", None),
                "color_b": _to_int(row[5] if len(row) > 5 else "", None),
                "fill_pattern_id": _to_int(row[6] if len(row) > 6 else "", None),
                "is_visible": _to_bool(row[7] if len(row) > 7 else "TRUE"),
            })
        return payload
    finally:
        try:
            if workbook is not None:
                workbook.Close(False)
        except Exception:
            pass
        try:
            if excel_app is not None:
                excel_app.Quit()
        except Exception:
            pass
        for obj in (used_range, worksheet, workbook, excel_app):
            try:
                if obj is not None:
                    Marshal.ReleaseComObject(obj)
            except Exception:
                pass


def build_entry_from_payload(item):
    storage_type = parse_storage_type(item.get("storage_type"))
    if storage_type is None:
        raise Exception("Unsupported storage type '{}'.".format(item.get("storage_type", "")))
    entry = DB.ColorFillSchemeEntry(storage_type)
    raw_value = item.get("value")
    if storage_type == DB.StorageType.String:
        value = "" if raw_value is None else str(raw_value)
    elif storage_type == DB.StorageType.Integer:
        value = _to_int(raw_value, 0)
    elif storage_type == DB.StorageType.Double:
        value = _to_float(raw_value, 0.0)
    elif storage_type == DB.StorageType.ElementId:
        value = DB.ElementId(_to_int(raw_value, -1))
    else:
        value = raw_value
    if not csu._entry_set_value(entry, storage_type, value):
        raise Exception("Failed to set value '{}' on entry.".format(raw_value))

    r = item.get("color_r")
    g = item.get("color_g")
    b = item.get("color_b")
    if r is not None and g is not None and b is not None and hasattr(entry, "Color"):
        try:
            entry.Color = DB.Color(int(r), int(g), int(b))
        except Exception:
            pass
    fill_pattern_id = item.get("fill_pattern_id")
    if fill_pattern_id is not None and hasattr(entry, "FillPatternId"):
        try:
            entry.FillPatternId = DB.ElementId(int(fill_pattern_id))
        except Exception:
            pass
    for attr in ("IsVisible", "Visible"):
        if hasattr(entry, attr):
            try:
                setattr(entry, attr, bool(item.get("is_visible", True)))
            except Exception:
                pass
    caption = item.get("caption", "")
    if caption:
        try:
            setter = getattr(entry, "SetCaption", None)
            if callable(setter):
                setter(caption)
            elif hasattr(entry, "Caption"):
                entry.Caption = caption
        except Exception:
            pass
    return entry


def apply_payload_to_scheme(target, payload, log=None):
    for attr, key in (("Title", "title"), ("IsByRange", "is_by_range"), ("IsByValue", "is_by_value"), ("IsByPercentage", "is_by_percentage")):
        if hasattr(target, attr) and key in payload:
            try:
                setattr(target, attr, payload.get(key))
            except Exception:
                pass

    entries_to_set = [build_entry_from_payload(item) for item in payload.get("entries", [])]
    set_entries = getattr(target, "SetEntries", None)
    if callable(set_entries):
        set_entries(csu._to_entry_collection(entries_to_set))
        csu._regenerate_scheme_document(target)
        refreshed = list(target.GetEntries())
        if refreshed:
            csu._patch_entry_colors(refreshed, entries_to_set, log=log, stage="import-post-set")
            set_entries(csu._to_entry_collection(refreshed))
            csu._regenerate_scheme_document(target)
        return True, None

    add_entry = getattr(target, "AddEntry", None)
    if not callable(add_entry):
        return False, "Target scheme API does not support SetEntries/AddEntry."
    failures = []
    for entry in entries_to_set:
        try:
            add_entry(entry)
        except Exception as ex:
            failures.append(str(ex))
    if failures:
        return False, "Failed to add one or more entries: {}".format(" | ".join(failures[:5]))
    csu._regenerate_scheme_document(target)
    return True, None


def get_scheme_definition_snapshot(doc, scheme):
    entry_storage_types = []
    try:
        entries = list(scheme.GetEntries())
    except Exception:
        entries = []
    for entry in entries:
        entry_storage_types.append(storage_type_token(getattr(entry, "StorageType", None)))
    unique_storage_types = sorted(set([s for s in entry_storage_types if s]))
    return {
        "scheme_name": getattr(scheme, "Name", "") or "Color Scheme",
        "category_id": elem_id_int(getattr(scheme, "CategoryId", None)),
        "category_name": category_name(doc, getattr(scheme, "CategoryId", None)),
        "area_scheme_name": scheme_area_scheme_name(doc, scheme),
        "title": getattr(scheme, "Title", "") or "",
        "parameter_id": _scheme_parameter_id_int(scheme),
        "parameter_name": _scheme_parameter_label(doc, scheme),
        "is_by_range": bool(getattr(scheme, "IsByRange", False)),
        "is_by_value": bool(getattr(scheme, "IsByValue", False)),
        "is_by_percentage": bool(getattr(scheme, "IsByPercentage", False)),
        "entry_storage_types": unique_storage_types,
    }


def validate_payload_against_scheme(doc, payload, target_scheme):
    target = get_scheme_definition_snapshot(doc, target_scheme)
    payload_storage_types = sorted(set([
        str(item.get("storage_type", "")).strip()
        for item in payload.get("entries", [])
        if str(item.get("storage_type", "")).strip()
    ]))
    problems = []
    if payload_storage_types and target.get("entry_storage_types") and payload_storage_types != target.get("entry_storage_types"):
        problems.append(
            "Entry storage type mismatch. Workbook={} Target={}".format(
                ", ".join(payload_storage_types),
                ", ".join(target.get("entry_storage_types") or []),
            )
        )
    payload_pid = payload.get("parameter_id")
    target_pid = target.get("parameter_id")
    if payload_pid is not None and target_pid is not None and payload_pid != target_pid:
        problems.append(
            "Scheme parameter mismatch. Workbook={} ({}) Target={} ({})".format(
                payload_pid,
                payload.get("parameter_name", ""),
                target_pid,
                target.get("parameter_name", ""),
            )
        )
    for key, label in (
        ("is_by_range", "By Range"),
        ("is_by_value", "By Value"),
        ("is_by_percentage", "By Percentage"),
    ):
        if key in payload and bool(payload.get(key)) != bool(target.get(key)):
            problems.append(
                "Scheme mode mismatch for {}. Workbook={} Target={}".format(
                    label,
                    bool(payload.get(key)),
                    bool(target.get(key)),
                )
            )
    ok = len(problems) == 0
    return ok, problems, target


def unique_scheme_name_in_scope(doc, scope_seed_scheme, base_name):
    base = (base_name or "Color Scheme").strip()
    if not base:
        base = "Color Scheme"
    all_schemes = csu.collect_color_fill_schemes(doc)
    if csu.find_scheme_in_scope_by_name(all_schemes, scope_seed_scheme, base) is None:
        return base
    idx = 1
    while idx < 1000:
        candidate = "{} ({})".format(base, idx)
        if csu.find_scheme_in_scope_by_name(all_schemes, scope_seed_scheme, candidate) is None:
            return candidate
        idx += 1
    return "{} ({})".format(base, "Copy")


def build_import_target_choices(doc, schemes):
    scope_map = {}
    for scheme in schemes:
        key = (elem_id_int(getattr(scheme, "CategoryId", None)), elem_id_int(scheme_area_scheme_id(scheme)))
        if key not in scope_map:
            scope_map[key] = scheme
    create_choices = [{
        "mode": "create",
        "label": "Create New In {}".format(scheme_scope_label(doc, seed)),
        "scheme": seed,
    } for seed in scope_map.values()]
    overwrite_choices = [{
        "mode": "overwrite",
        "label": "Overwrite {}".format(scheme_display_name(doc, scheme)),
        "scheme": scheme,
    } for scheme in schemes]
    create_choices.sort(key=lambda x: x["label"].lower())
    overwrite_choices.sort(key=lambda x: x["label"].lower())
    return create_choices + overwrite_choices


def choose_default_import_target_index(doc, payload, choices):
    payload_area = (payload.get("area_scheme_name", "") or "").strip().lower()
    payload_category = (payload.get("category_name", "") or "").strip().lower()
    payload_scheme = (payload.get("scheme_name", "") or "").strip().lower()
    for idx, choice in enumerate(choices):
        scheme = choice.get("scheme")
        if not scheme:
            continue
        area_name = scheme_area_scheme_name(doc, scheme).strip().lower()
        cat_name = category_name(doc, getattr(scheme, "CategoryId", None)).strip().lower()
        if payload_area:
            if area_name != payload_area:
                continue
        elif payload_category and cat_name != payload_category:
            continue
        if choice.get("mode") == "overwrite":
            name = (getattr(scheme, "Name", "") or "").strip().lower()
            if name == payload_scheme:
                return idx
        else:
            return idx
    return 0
